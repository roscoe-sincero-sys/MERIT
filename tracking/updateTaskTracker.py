"""
Builds tracking/MERIT_Task_Tracker.xlsx from tracking/tasks.json.

tasks.json is the source of truth: edit a task's status there (Not started, In progress, Done,
Deferred, Blocked) and re-run this script. The workbook gets a Tasks sheet and a Summary sheet
that rolls progress up by phase.

    python tracking/updateTaskTracker.py
"""
from __future__ import annotations

import json
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "tasks.json"
WORKBOOK = HERE / "MERIT_Task_Tracker.xlsx"

STATUS_COLORS = {
    "Done": "C6EFCE",
    "In progress": "FFEB9C",
    "Not started": "FFFFFF",
    "Deferred": "D9D9D9",
    "Blocked": "FFC7CE",
}
HEADERS = ["ID", "Phase", "Task", "Spec reference", "Deliverable / location", "Status",
           "% complete", "Started", "Finished", "Notes"]


def main() -> int:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    tasks = data["tasks"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    ws.append(HEADERS)
    for t in tasks:
        pct = t.get("pct")
        if pct is None:
            pct = 100 if t["status"] == "Done" else 0
        ws.append([t["id"], t["phase"], t["task"], t.get("ref", ""), t.get("where", ""), t["status"],
                   pct / 100, t.get("started", ""), t.get("finished", ""), t.get("notes", "")])

    fill = PatternFill("solid", fgColor="1A4480")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate([8, 24, 60, 30, 40, 13, 11, 12, 12, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[6].number_format = "0%"
        for c in (row[2], row[3], row[4], row[9]):
            c.alignment = Alignment(wrap_text=True, vertical="top")
    last = ws.max_row
    for status, color in STATUS_COLORS.items():
        ws.conditional_formatting.add(
            f"F2:F{last}", CellIsRule(operator="equal", formula=[f'"{status}"'],
                                      fill=PatternFill("solid", fgColor=color)))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary", 0)
    summary.append(["MERIT development progress"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} from tracking/tasks.json"])
    summary.append([])
    summary.append(["Phase", "Tasks", "Done", "In progress", "Not started", "Deferred", "Blocked",
                    "% complete (excl. deferred)"])
    for c in summary[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    phases: "OrderedDict[str, list]" = OrderedDict()
    for t in tasks:
        phases.setdefault(t["phase"], []).append(t)

    def row_for(name, items):
        counts = {s: sum(1 for t in items if t["status"] == s) for s in STATUS_COLORS}
        active = [t for t in items if t["status"] != "Deferred"]
        pct = 0.0
        if active:
            pct = sum((t.get("pct") if t.get("pct") is not None else (100 if t["status"] == "Done" else 0))
                      for t in active) / (100 * len(active))
        return [name, len(items), counts["Done"], counts["In progress"], counts["Not started"],
                counts["Deferred"], counts["Blocked"], pct]

    for name, items in phases.items():
        summary.append(row_for(name, items))
    summary.append(row_for("ALL", tasks))
    for c in summary[summary.max_row]:
        c.font = Font(bold=True)
    for row in summary.iter_rows(min_row=5):
        row[7].number_format = "0%"
    for i, w in enumerate([34, 8, 8, 12, 12, 10, 9, 24], start=1):
        summary.column_dimensions[get_column_letter(i)].width = w

    wb.save(WORKBOOK)
    print(f"Wrote {WORKBOOK} ({len(tasks)} tasks)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
