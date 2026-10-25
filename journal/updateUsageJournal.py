"""
Builds journal/MERIT_ClaudeCode_Usage.xlsx from Claude Code's local session transcripts.

Claude Code writes one JSONL transcript per session under
    %USERPROFILE%/.claude/projects/<project-slug>/<sessionId>.jsonl
Each assistant message carries a "usage" block (input, cache-creation, cache-read and output
tokens) and every entry carries a UTC timestamp. This script totals them per session, so the
start time, end time, duration and token use of every session on the MERIT project are recorded.

The /usage percentage shown inside Claude Code is not written to the transcript. The sheet keeps a
"Plan usage % (from /usage)" column for it; values typed there by hand are preserved on re-run.

Usage:
    python journal/updateUsageJournal.py            # re-scan transcripts and rewrite the sheet
    python journal/updateUsageJournal.py --note "<sessionId>=<text>"
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

JOURNAL_DIR = Path(__file__).resolve().parent
PROJECT_DIR = JOURNAL_DIR.parent
WORKBOOK = JOURNAL_DIR / "MERIT_ClaudeCode_Usage.xlsx"


def project_slug(path: Path) -> str:
    # Claude Code replaces ':', '\' and '/' in the project path with '-'.
    text = str(path)
    for ch in (":", "\\", "/"):
        text = text.replace(ch, "-")
    return text


def transcript_dir() -> Path:
    return Path.home() / ".claude" / "projects" / project_slug(PROJECT_DIR)


def parse_ts(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def summarize(path: Path) -> dict:
    first = last = None
    model = None
    version = None
    prompts = 0
    turns = 0
    tok = {"input": 0, "cacheCreate": 0, "cacheRead": 0, "output": 0}
    seen_ids: set[str] = set()
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            ts = row.get("timestamp")
            if ts:
                t = parse_ts(ts)
                first = t if first is None or t < first else first
                last = t if last is None or t > last else last
            version = row.get("version") or version
            if row.get("type") == "user" and row.get("origin", {}).get("kind") == "human":
                prompts += 1
            msg = row.get("message") or {}
            if row.get("type") == "assistant" and isinstance(msg, dict):
                model = msg.get("model") or model
                usage = msg.get("usage")
                # One API response can be split across several transcript lines; count it once.
                mid = msg.get("id")
                if usage and mid not in seen_ids:
                    if mid:
                        seen_ids.add(mid)
                    turns += 1
                    tok["input"] += usage.get("input_tokens", 0) or 0
                    tok["cacheCreate"] += usage.get("cache_creation_input_tokens", 0) or 0
                    tok["cacheRead"] += usage.get("cache_read_input_tokens", 0) or 0
                    tok["output"] += usage.get("output_tokens", 0) or 0
    return {
        "sessionId": path.stem,
        "model": model or "",
        "version": version or "",
        "start": first,
        "end": last,
        "prompts": prompts,
        "turns": turns,
        **tok,
    }


HEADERS = [
    "Session ID", "Model", "Claude Code version", "Start (local)", "End (local)", "Duration (h:mm)",
    "Human prompts", "API turns", "Input tokens", "Cache-write tokens", "Cache-read tokens",
    "Output tokens", "Total tokens", "Plan usage % (from /usage)", "Notes",
]
MANUAL_COLS = {"Plan usage % (from /usage)", "Notes"}


def existing_manual_values() -> dict[str, dict[str, object]]:
    if not WORKBOOK.exists():
        return {}
    wb = load_workbook(WORKBOOK)
    ws = wb["Sessions"]
    header = [c.value for c in ws[1]]
    out: dict[str, dict[str, object]] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        rec = dict(zip(header, r))
        out[str(r[0])] = {k: rec.get(k) for k in MANUAL_COLS}
    return out


def local(dt: datetime | None):
    return dt.astimezone().replace(tzinfo=None) if dt else None


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--note", action="append", default=[], help="<sessionId>=<text>")
    ap.add_argument("--usage", action="append", default=[], help="<sessionId>=<percent>")
    args = ap.parse_args()

    tdir = transcript_dir()
    if not tdir.exists():
        print(f"No transcripts found at {tdir}", file=sys.stderr)
        return 1

    manual = existing_manual_values()
    for item in args.note:
        sid, _, text = item.partition("=")
        manual.setdefault(sid, {})["Notes"] = text
    for item in args.usage:
        sid, _, pct = item.partition("=")
        manual.setdefault(sid, {})["Plan usage % (from /usage)"] = pct

    sessions = [summarize(p) for p in sorted(tdir.glob("*.jsonl"))]
    sessions = [s for s in sessions if s["start"]]
    sessions.sort(key=lambda s: s["start"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.append(HEADERS)
    for s in sessions:
        dur = (s["end"] - s["start"]) if s["end"] else None
        dur_txt = ""
        if dur:
            mins = int(dur.total_seconds() // 60)
            dur_txt = f"{mins // 60}:{mins % 60:02d}"
        total = s["input"] + s["cacheCreate"] + s["cacheRead"] + s["output"]
        m = manual.get(s["sessionId"], {})
        ws.append([
            s["sessionId"], s["model"], s["version"], local(s["start"]), local(s["end"]), dur_txt,
            s["prompts"], s["turns"], s["input"], s["cacheCreate"], s["cacheRead"], s["output"], total,
            m.get("Plan usage % (from /usage)"), m.get("Notes"),
        ])

    head_fill = PatternFill("solid", fgColor="1A4480")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [38, 18, 12, 19, 19, 10, 9, 9, 12, 14, 15, 12, 14, 14, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = row[4].number_format = "yyyy-mm-dd hh:mm:ss"
        for c in row[8:13]:
            c.number_format = "#,##0"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    n = len(sessions) + 1
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", "", f"=SUM(G2:G{n})", f"=SUM(H2:H{n})", f"=SUM(I2:I{n})",
               f"=SUM(J2:J{n})", f"=SUM(K2:K{n})", f"=SUM(L2:L{n})", f"=SUM(M2:M{n})"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.number_format = "#,##0"

    about = wb.create_sheet("About")
    about.column_dimensions["A"].width = 110
    for line in [
        "MERIT - Claude Code session usage journal",
        f"Last refreshed: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} (local time)",
        f"Source: {tdir}",
        "",
        "Refresh: python journal\\updateUsageJournal.py",
        "Record a /usage reading: python journal\\updateUsageJournal.py --usage \"<sessionId>=42%\"",
        "Times are converted from the transcript's UTC timestamps to this computer's local time zone.",
        "Start = first transcript entry, End = last transcript entry (the session may still be open).",
        "Cache-read tokens are billed at a fraction of input tokens; Total is a raw sum, not a cost.",
        "Stop rule agreed for this project: work continues until plan usage reaches 100%.",
    ]:
        about.append([line])
    about["A1"].font = Font(bold=True, size=14)

    wb.save(WORKBOOK)
    print(f"Wrote {WORKBOOK} ({len(sessions)} session(s))")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
